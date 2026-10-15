"""Сервис расписания на основе Excel-файла schedule.xlsx.

Парсит лист «Слоты» и предоставляет методы для работы со слотами:
    - get_free_slots()     — список свободных слотов по фильтрам;
    - is_slot_available()  — проверка конкретного слота;
    - book_slot()          — пометить слот занятым в памяти;
    - resolve_doctor_id()  — маппинг врача из Excel в ID справочника MIS.

Особенности:
    - «Сегодня» читается из settings.SCHEDULE_TODAY (по умолчанию 14.10.2026),
      потому что данные файла относятся к 2026 году. В проде — datetime.now().
    - Слоты со статусом «Свободен» и «Отменён пациентом» считаются доступными.
    - ФИО врача сопоставляется с справочником MIS; при отсутствии совпадения
      используется первый активный врач той же специальности.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

from app.config import settings
from app.core.exceptions import MISNotFound
from app.services.mis_client import MISClientProtocol


FREE_STATUSES = {"Свободен", "Отменён пациентом"}
SHEET_NAME = "Слоты"

# Колонки листа «Слоты» (1-based индексы для openpyxl)
COL_ID = 1
COL_DATE = 2
COL_DOCTOR_ID = 4
COL_DOCTOR_NAME = 5
COL_SPECIALTY = 6
COL_START = 9
COL_STATUS = 12


# Модель слота
@dataclass
class Slot:
    """Один слот приёма из расписания."""

    slot_id: str
    date: str          # ДД.ММ.ГГГГ
    doctor_excel_id: str  # D001
    doctor_name: str
    specialty: str
    start: str         # ЧЧ:ММ
    status: str

    @property
    def scheduled_at(self) -> str:
        """Формат для MIS: «ДД.ММ.ГГГГ ЧЧ:ММ»."""
        return f"{self.date} {self.start}"

    @property
    def is_free(self) -> bool:
        return self.status in FREE_STATUSES



# Сервис
class ScheduleService:
    """Сервис расписания. Загружает слоты из Excel в память."""

    def __init__(self, xlsx_path: str | None = None) -> None:
        self._path = Path(xlsx_path or settings.SCHEDULE_XLSX_PATH)
        self._slots: list[Slot] = []
        self._booked: set[tuple[str, str]] = set()  # (doctor_excel_id, scheduled_at)
        self._today = self._resolve_today()
        self._loaded = False

   
    # Загрузка
    def _resolve_today(self) -> str:
        """Определить «сегодня». Если SCHEDULE_TODAY задан — берём его."""
        raw = (settings.SCHEDULE_TODAY or "").strip()
        if raw:
            return raw
        return datetime.now().strftime("%d.%m.%Y")

    def load(self) -> None:
        """Загрузить лист «Слоты» из Excel. Идемпотентно."""
        if self._loaded:
            return

        if not self._path.exists():
            raise FileNotFoundError(
                f"Файл расписания не найден: {self._path}. "
                f"Скачай schedule.xlsx из материалов кейса и положи в data/."
            )

        # Импорт здесь, чтобы не тянуть openpyxl, если schedule не нужен
        from openpyxl import load_workbook

        wb = load_workbook(self._path, read_only=True, data_only=True)
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"В файле {self._path} нет листа «{SHEET_NAME}». "
                f"Доступные листы: {wb.sheetnames}"
            )

        ws = wb[SHEET_NAME]
        slots: list[Slot] = []

        # Пропускаем шапку (1-я строка)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[COL_ID - 1] is None:
                continue

            slot_id = str(row[COL_ID - 1]).strip()
            date_raw = row[COL_DATE - 1]
            doctor_excel_id = str(row[COL_DOCTOR_ID - 1]).strip()
            doctor_name = str(row[COL_DOCTOR_NAME - 1]).strip()
            specialty = str(row[COL_SPECIALTY - 1]).strip()
            start_raw = row[COL_START - 1]
            status = str(row[COL_STATUS - 1]).strip()

            date_str = self._format_date(date_raw)
            start_str = self._format_time(start_raw)

            if not date_str or not start_str:
                continue

            slots.append(
                Slot(
                    slot_id=slot_id,
                    date=date_str,
                    doctor_excel_id=doctor_excel_id,
                    doctor_name=doctor_name,
                    specialty=specialty,
                    start=start_str,
                    status=status,
                )
            )

        wb.close()
        self._slots = slots
        self._loaded = True

    @staticmethod
    def _format_date(value: object) -> str:
        """Привести дату из Excel к «ДД.ММ.ГГГГ»."""
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%d.%m.%Y")
        text = str(value).strip()
        # Если уже в формате ДД.ММ.ГГГГ
        try:
            datetime.strptime(text, "%d.%m.%Y")
            return text
        except ValueError:
            pass
        # Если ISO
        try:
            return datetime.fromisoformat(text).strftime("%d.%m.%Y")
        except ValueError:
            return ""

    @staticmethod
    def _format_time(value: object) -> str:
        """Привести время из Excel к «ЧЧ:ММ»."""
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%H:%M")
        text = str(value).strip()
        # «14:00» или «14:00:00»
        parts = text.split(":")
        if len(parts) >= 2:
            try:
                hh = int(parts[0])
                mm = int(parts[1])
                return f"{hh:02d}:{mm:02d}"
            except ValueError:
                return ""
        return ""

   
    # Поиск слотов
    async def get_free_slots(
        self,
        specialty: str | None = None,
        doctor_excel_id: str | None = None,
        date_from: str | None = None,
        date_to: str | None = None,
        limit: int = 5,
    ) -> list[Slot]:
        """Свободные слоты на будущие даты.

        Аргументы:
            specialty: фильтр по специальности («Терапевт»).
            doctor_excel_id: фильтр по конкретному врачу («D001»).
            date_from: нижняя граница, ДД.ММ.ГГГГ. По умолчанию — сегодня.
            date_to: верхняя граница, ДД.ММ.ГГГГ. По умолчанию — +30 дней.
            limit: сколько слотов вернуть.
        """
        self.load()

        today = self._today
        if not date_from:
            date_from = today
        if not date_to:
            dt = datetime.strptime(today, "%d.%m.%Y") + timedelta(days=30)
            date_to = dt.strftime("%d.%m.%Y")

        needle_spec = specialty.strip().lower() if specialty else None
        doctor = doctor_excel_id.strip() if doctor_excel_id else None

        result: list[Slot] = []
        for slot in self._slots:
            if not slot.is_free:
                continue
            if self._is_booked(slot):
                continue
            if needle_spec and slot.specialty.strip().lower() != needle_spec:
                continue
            if doctor and slot.doctor_excel_id != doctor:
                continue
            if not self._date_in_range(slot.date, date_from, date_to):
                continue

            result.append(slot)
            if len(result) >= limit:
                break

        return result

    def _date_in_range(self, date_str: str, date_from: str, date_to: str) -> bool:
        try:
            d = datetime.strptime(date_str, "%d.%m.%Y")
            df = datetime.strptime(date_from, "%d.%m.%Y")
            dt = datetime.strptime(date_to, "%d.%m.%Y")
        except ValueError:
            return False
        return df <= d <= dt

    def _is_booked(self, slot: Slot) -> bool:
        return (slot.doctor_excel_id, slot.scheduled_at) in self._booked

    async def is_slot_available(
        self,
        doctor_excel_id: str,
        scheduled_at: str,
    ) -> bool:
        """Свободен ли конкретный слот (по ID врача из Excel и дате-времени)."""
        self.load()
        for slot in self._slots:
            if slot.doctor_excel_id != doctor_excel_id:
                continue
            if slot.scheduled_at != scheduled_at:
                continue
            return slot.is_free and not self._is_booked(slot)
        return False

    async def book_slot(self, doctor_excel_id: str, scheduled_at: str) -> None:
        """Пометить слот занятым (in-memory, для MVP)."""
        self._booked.add((doctor_excel_id, scheduled_at))

    def unbook_slot(self, doctor_excel_id: str, scheduled_at: str) -> None:
        """Освободить слот (например, при отмене)."""
        self._booked.discard((doctor_excel_id, scheduled_at))

    
    # Маппинг врачей
    async def resolve_doctor_id(
        self,
        doctor_excel_id: str,
        mis: MISClientProtocol,
    ) -> int:
        """Сопоставить врача из Excel с ID в справочнике MIS.

        1. Ищем врача по ФИО среди активных.
        2. Если не нашли — берём первого активного той же специальности.
        3. Если и таких нет — MISNotFound.
        """
        self.load()

        # Находим слота врача в Excel, чтобы узнать ФИО и специальность
        doctor_name: str | None = None
        specialty: str | None = None
        for slot in self._slots:
            if slot.doctor_excel_id == doctor_excel_id:
                doctor_name = slot.doctor_name
                specialty = slot.specialty
                break

        if doctor_name is None:
            raise MISNotFound(
                f"Врач {doctor_excel_id} не найден в расписании Excel",
                code="НЕ_НАЙДЕНО",
                status_code=404,
            )

        # 1. Пробуем по ФИО
        candidates = await mis.list_doctors(specialty=specialty, limit=50)
        for d in candidates:
            if d.is_active and d.full_name.strip() == doctor_name.strip():
                return d.id

        # 2. Fallback — первый активный той же специальности
        for d in candidates:
            if d.is_active:
                return d.id

        raise MISNotFound(
            f"Нет активного врача специальности «{specialty}» в справочнике MIS",
            code="НЕ_НАЙДЕНО",
            status_code=404,
        )


# Синглтон
_schedule: ScheduleService | None = None


def get_schedule() -> ScheduleService:
    """Получить глобальный экземпляр ScheduleService."""
    global _schedule
    if _schedule is None:
        _schedule = ScheduleService()
    return _schedule